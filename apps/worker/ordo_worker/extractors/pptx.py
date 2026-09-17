from io import BytesIO

from pptx import Presentation


def extract_pptx_text(content: bytes) -> str:
    presentation = Presentation(BytesIO(content))
    lines: list[str] = []
    for slide in presentation.slides:
        for shape in slide.shapes:
            if not shape.has_text_frame:
                continue
            for paragraph in shape.text_frame.paragraphs:
                text = "".join(run.text for run in paragraph.runs)
                if text.strip():
                    lines.append(text)
    return "\n".join(lines)
